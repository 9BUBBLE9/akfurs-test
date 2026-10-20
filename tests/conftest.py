import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

CRM_HEADER = ("Номер заказа;№ строки;Дата создания;Клиент;Менеджер;Источник;Артикул;Кол-во;"
              "Цена за ед.;Валюта;Скидка;Статус;Комментарий;Выгружено")


def write_catalog(path, renames=None, prices=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Каталог"
    ws.append(["Учебный каталог"])  # строка-заголовок над таблицей, как в оригинале
    ws.append(["Артикул", "Наименование"])
    for article in ["OLD-01", "NEW01N", "KEEP-01"]:
        ws.append([article, "изделие"])

    ws = wb.create_sheet("Переименования")
    ws.append(["Старый артикул", "Новый артикул", "Дата перехода"])
    for row in renames if renames is not None else [("OLD-01", "NEW01N", dt.datetime(2025, 6, 1))]:
        ws.append(list(row))

    ws = wb.create_sheet("Прайс_история")
    ws.sheet_state = "hidden"
    ws.append(["Артикул", "Действует с", "Действует по",
               "Рекомендованная цена, ₽", "Минимально допустимая цена, ₽"])
    default_prices = [
        ("NEW01N", dt.datetime(2025, 1, 1), dt.datetime(2025, 12, 31), 120_000, 100_000),
        ("NEW01N", dt.datetime(2026, 1, 1), None, 130_000, 110_000),
        ("KEEP-01", dt.datetime(2025, 1, 1), None, 60_000, 50_000),
    ]
    for row in prices if prices is not None else default_prices:
        ws.append(list(row))

    ws = wb.create_sheet("Менеджеры")
    ws.append(["Каноническое имя", "Варианты написания в CRM", "Дата выхода", "Дата увольнения"])
    ws.append(["Ирина Зайцева", "Ирина Зайцева | Зайцева И. | Irina Z | И. Зайцева",
               dt.datetime(2025, 1, 1), None])
    ws.append(["Дарья Белкина", "Дарья Белкина | darya.b | Дарья", dt.datetime(2025, 1, 1), None])
    wb.save(path)


def write_crm(path, rows, encoding="utf-8-sig"):
    Path(path).write_text("\n".join([CRM_HEADER, *rows]) + "\n", encoding=encoding)


@pytest.fixture
def data_dir(tmp_path):
    write_catalog(tmp_path / "03_catalog.xlsx")
    write_crm(tmp_path / "01_crm_orders_export.csv", [
        # дубль строки 1 заказа 7: старая версия ниже в файле, новая выше
        "A-00007;1;2030-01-10;CL-90001; И. Зайцева;IG;OLD-01;1;150 000;руб;10%;PAID;;2030-01-01 09:00",
        "А-7;2;10.01.2030;cl90001;Irina Z;Insta;KEEP-01;2;30000,00;₽;5 000;оплачено;;2030-01-01 09:00",
        "a-7;1;10/01/2030;КЛ-90001;Зайцева И.;Instagram;OLD-01;1;140 000;RUB;;NEW;;2029-12-01 12:00",
        "A00008;1;45667;CL 90002;darya.b;Сарафан;NEW01N;1;98,000.00;RUR;;отменен;;2030-01-01 09:00",
    ])
    return tmp_path
